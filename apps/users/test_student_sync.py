import io
from pathlib import Path
import random
import string
import tempfile
from datetime import time
from unittest.mock import patch

from django.conf import settings
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import SimpleTestCase, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from apps.nightpass.models import CampusResource, Hostel
from apps.users.models import CustomUser, NightPass, Student, ViolationAuditLog
from apps.users.services.student_sync import apply_sync, preview_sync, read_upload


class StudentSyncTests(TestCase):
    def setUp(self):
        self.hostel = Hostel.objects.create(name="A", email="a@example.com", contact_number="123")
        self.other_hostel = Hostel.objects.create(name="B", email="b@example.com", contact_number="123")
        self.user = CustomUser.objects.create_user("Student@Example.com", "password")
        self.student = Student.objects.create(user=self.user, registration_number="001", name="Original",
                                              email="Student@Example.com", hostel=self.hostel, room_number="101",
                                              picture="https://example.com/old.jpg", gender="male", violation_flags=3,
                                              last_scan_at=timezone.now(), has_booked=True, is_checked_in=False)

    def sync(self, rows, mode="full", **kwargs):
        return apply_sync(list(rows[0]), rows, mode, actor="test-admin", filename="test.csv", **kwargs)

    def snapshot(self):
        return (list(Student.objects.order_by("pk").values()), list(CustomUser.objects.order_by("pk").values()))

    def test_existing_student_case_insensitive_email(self):
        self.sync([{"email": " STUDENT@example.COM ", "hostel": "B"}])
        self.student.refresh_from_db()
        self.assertEqual(self.student.hostel, self.other_hostel)
        self.assertEqual(self.student.email, "student@example.com")
        self.assertEqual(self.student.name, "Original")
        self.assertEqual(CustomUser.objects.count(), 1)

    def test_reuse_existing_user(self):
        user = CustomUser.objects.create_user("Reuse@Example.com", "password")
        before = CustomUser.objects.get(pk=user.pk).__dict__.copy()
        counts, _ = self.sync([{"email": " reuse@example.com ", "name": "Reuse", "registration_number": "002"}])
        self.assertEqual(Student.objects.get(pk="002").user_id, user.pk)
        self.assertEqual(counts["users_to_reuse"], 1)
        user.refresh_from_db()
        self.assertEqual(user.password, before["password"])
        self.assertEqual(user.email, before["email"])

    def test_create_new_user_with_unusable_password_and_normalized_gender(self):
        counts, _ = self.sync([{"email": " New@Example.com ", "name": "New", "registration_number": "002", "gender": "FEMALE"}])
        student = Student.objects.get(pk="002")
        self.assertEqual(student.gender, "female")
        self.assertEqual(student.user.email, "new@example.com")
        self.assertEqual(student.user.user_type, "student")
        self.assertFalse(student.user.has_usable_password())
        self.assertEqual(counts["users_to_create"], 1)

    def test_picture_mode_only_writes_picture_and_never_downloads(self):
        before = Student.objects.values().get(pk="001")
        with patch("requests.get", side_effect=AssertionError("no downloads")), patch.object(Student, "save", side_effect=AssertionError("no per-row saves")):
            self.sync([{"email": "student@example.com", "picture": "https://example.com/new.jpg", "name": "Ignored", "hostel": "Unknown"}], "picture")
        after = Student.objects.values().get(pk="001")
        before["picture"] = "https://example.com/new.jpg"
        self.assertEqual(after, before)

    def test_blank_picture_default_preserves_and_explicit_option_clears(self):
        row = [{"email": "student@example.com", "picture": ""}]
        counts, _ = self.sync(row, "picture")
        self.student.refresh_from_db()
        self.assertTrue(self.student.picture)
        self.assertEqual(counts["skipped_rows"], 1)
        self.sync(row, "picture", allow_blank_picture=True)
        self.student.refresh_from_db()
        self.assertIsNone(self.student.picture)

    def test_hostel_mode_only_changes_assignments(self):
        before = Student.objects.values().get(pk="001")
        self.sync([{"email": "student@example.com", "hostel": "B", "room_number": "202", "name": "Ignored"}], "hostel")
        before.update(hostel_id="B", room_number="202")
        self.assertEqual(Student.objects.values().get(pk="001"), before)

    def test_clear_all_is_one_update_and_restores_unchanged_assignments(self):
        second = Student.objects.create(user=CustomUser.objects.create_user("second@example.com", None), registration_number="002", name="Second", hostel=self.hostel, room_number="102")
        with CaptureQueriesContext(connection) as queries:
            self.sync([{"email": "student@example.com", "hostel": "A", "room_number": "101"}], "hostel", clear="all")
        clear_queries = [q["sql"] for q in queries if q["sql"].startswith('UPDATE "users_student"') and 'CASE' not in q["sql"]]
        self.assertEqual(len(clear_queries), 1)
        second.refresh_from_db()
        self.student.refresh_from_db()
        self.assertIsNone(second.hostel)
        self.assertIsNone(second.room_number)
        self.assertEqual(self.student.hostel_id, "A")
        self.assertEqual(self.student.room_number, "101")

    def test_clear_absent(self):
        Student.objects.create(user=CustomUser.objects.create_user("second@example.com", None), registration_number="002", name="Second", hostel=self.hostel)
        counts, _ = self.sync([{"email": "student@example.com", "hostel": "B", "room_number": "202"}], "hostel", clear="absent")
        self.assertEqual(counts["assignments_to_clear"], 1)
        self.assertIsNone(Student.objects.get(pk="002").hostel_id)
        self.assertEqual(Student.objects.get(pk="001").hostel_id, "B")

    def test_validation_errors_never_write_even_with_clear_all(self):
        for row in (
            {"email": "student@example.com", "hostel": "Unknown", "room_number": "202"},
            {"email": "", "name": "Malformed Student", "hostel": "B", "room_number": "202"},
            {"email": "missing@example.com", "hostel": "B", "room_number": "202"},
        ):
            before = self.snapshot()
            with self.subTest(row=row), self.assertRaises(ValidationError):
                self.sync([row], "hostel", clear="all")
            self.assertEqual(self.snapshot(), before)

    def test_protected_accounts(self):
        for i, flags in enumerate(({"user_type": "admin"}, {"user_type": "security"}, {"is_staff": True}, {"is_superuser": True})):
            user = CustomUser.objects.create_user(f"protected{i}@example.com", None, **flags)
            before = self.snapshot()
            with self.subTest(flags=flags), self.assertRaises(ValidationError):
                self.sync([{"email": user.email.upper(), "name": "Overwrite", "registration_number": f"P{i}"}])
            self.assertEqual(before, self.snapshot())

    def test_duplicate_email_detection(self):
        rows = [{"email": "student@example.com", "name": "One"}, {"email": " STUDENT@EXAMPLE.COM ", "name": "Two"}]
        plan = preview_sync(list(rows[0]), rows, "full")
        self.assertEqual(plan.counts["duplicate_emails"], 1)
        self.assertEqual(plan.counts["error_rows"], 2)
        with self.assertRaises(ValidationError):
            self.sync(rows)

    def test_registration_conflicts_and_primary_key_changes(self):
        cases = [
            [{"email": "new@example.com", "registration_number": "001", "name": "New"}],
            [{"email": "new@example.com", "registration_number": "002", "name": "New"}, {"email": "other@example.com", "registration_number": "002", "name": "Other"}],
            [{"email": "student@example.com", "registration_number": "999"}],
        ]
        for rows in cases:
            before = self.snapshot()
            with self.subTest(rows=rows), self.assertRaises(ValidationError):
                self.sync(rows)
            self.assertEqual(before, self.snapshot())

    def test_invalid_later_row_no_partial_creation(self):
        before = self.snapshot()
        with self.assertRaises(ValidationError):
            self.sync([{"email": "new@example.com", "registration_number": "002", "name": "New"}, {"email": "invalid", "registration_number": "003", "name": "Bad"}])
        self.assertEqual(before, self.snapshot())

    def test_exception_rolls_back_users_and_hostel_clear(self):
        before = self.snapshot()
        with patch.object(Student.objects, "bulk_create", side_effect=RuntimeError("simulated failure")), self.assertRaises(RuntimeError):
            self.sync([{"email": "new@example.com", "registration_number": "002", "name": "New"}])
        self.assertEqual(before, self.snapshot())
        with patch.object(Student.objects, "bulk_update", side_effect=RuntimeError("simulated failure")), self.assertRaises(RuntimeError):
            self.sync([{"email": "student@example.com", "hostel": "B", "room_number": "202"}], "hostel", clear="all")
        self.assertEqual(before, self.snapshot())

    def test_preview_is_read_only(self):
        rows = [{"email": "new@example.com", "registration_number": "002", "name": "New"}]
        with CaptureQueriesContext(connection) as queries:
            plan = preview_sync(list(rows[0]), rows, "full")
        self.assertTrue(all(q["sql"].startswith("SELECT") for q in queries))
        self.assertEqual(plan.counts["users_to_create"], 1)

    def test_passes_violations_and_scan_fields_untouched(self):
        resource = CampusResource.objects.create(name="Library", max_capacity=10, start_time=time(0), end_time=time(23))
        night_pass = NightPass.objects.create(user=self.user, campus_resource=resource, start_time=time(20), end_time=timezone.now())
        ViolationAuditLog.objects.create(student=self.student, night_pass=night_pass, event_type="BECAME_DEFAULTER", message="Existing")
        passes = list(NightPass.objects.values())
        audits = list(ViolationAuditLog.objects.values())
        before = Student.objects.values().get(pk="001")
        for mode, row in (("full", {"email": "student@example.com", "name": "Changed"}), ("hostel", {"email": "student@example.com", "hostel": "B", "room_number": "202"}), ("picture", {"email": "student@example.com", "picture": "https://example.com/new.jpg"})):
            self.sync([row], mode)
        after = Student.objects.values().get(pk="001")
        for key in set(before) - {"name", "email", "hostel_id", "room_number", "picture"}:
            self.assertEqual(before[key], after[key], key)
        self.assertEqual(passes, list(NightPass.objects.values()))
        self.assertEqual(audits, list(ViolationAuditLog.objects.values()))

    def test_ambiguous_existing_user_emails_block(self):
        CustomUser.objects.create_user("STUDENT@example.com", None)
        with self.assertRaises(ValidationError):
            self.sync([{"email": "student@example.com", "name": "Changed"}])

    def test_mismatched_profile_email_blocks(self):
        Student.objects.filter(pk="001").update(email="different@example.com")
        with self.assertRaises(ValidationError):
            self.sync([{"email": "different@example.com", "name": "Changed"}])

    def test_csv_and_xlsx(self):
        headers, rows = read_upload(SimpleUploadedFile("students.csv", b'\xef\xbb\xbfemail,name,registration_number\r\nnew@example.com,"Last, First",002\r\n'))
        self.assertEqual(rows[0]["name"], "Last, First")
        self.assertEqual(rows[0]["registration_number"], "002")
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(headers)
        book.active.append(["new@example.com", "Name", "002"])
        output = io.BytesIO()
        book.save(output)
        _, rows = read_upload(SimpleUploadedFile("students.xlsx", output.getvalue()))
        self.assertEqual(rows[0]["registration_number"], "002")

    def test_bulk_query_count(self):
        rows = [{"email": f"bulk{i}@example.com", "registration_number": f"B{i}", "name": "Bulk"} for i in range(1000)]
        with CaptureQueriesContext(connection) as queries:
            self.sync(rows)
        self.assertLess(len(queries), 100)
        self.assertEqual(Student.objects.count(), 1001)

    def test_invalid_field_lengths_gender_and_url(self):
        for values in ({"name": "x" * 101}, {"gender": "unknown"}, {"picture": "not-a-url"}, {"room_number": "x" * 11}):
            before = self.snapshot()
            with self.subTest(values=values), self.assertRaises(ValidationError):
                self.sync([{"email": "student@example.com", **values}])
            self.assertEqual(before, self.snapshot())

    def test_missing_headers_and_malformed_csv(self):
        for data in (b"email,email\na@example.com,b@example.com\n", b"email,name\na@example.com\n", b""):
            with self.subTest(data=data), self.assertRaises(ValidationError):
                read_upload(SimpleUploadedFile("students.csv", data))
        with self.assertRaises(ValidationError):
            preview_sync(["name"], [{"name": "New"}], "full")

    def test_duplicate_existing_student_profile_emails(self):
        user = CustomUser.objects.create_user("other@example.com", None)
        Student.objects.create(user=user, registration_number="002", name="Other", email="STUDENT@example.com")
        with self.assertRaises(ValidationError):
            self.sync([{"email": "student@example.com", "name": "Changed"}])


class StudentSyncAdminTests(TestCase):
    REAL_HEADERS = "registration_number,name,hostel,Caretaker Name,gender,room_number,contact_number,email,parent_contact,year,user,URL"

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.storage = Path(directory.name)
        storage_settings = override_settings(STUDENT_SYNC_TEMP_DIR=directory.name)
        storage_settings.enable()
        self.addCleanup(storage_settings.disable)
        self.admin = CustomUser.objects.create_superuser("admin@example.com", "password")
        self.client.force_login(self.admin)
        self.url = reverse("admin:users_student_data_sync").removeprefix(settings.FORCE_SCRIPT_NAME or "")

    def upload(self, mode="full", data=b"email,registration_number,name\nnew@example.com,001,New\n"):
        return self.client.post(self.url, {"mode": mode, "file": SimpleUploadedFile("students.csv", data)})

    def test_admin_preview_apply_and_subpath(self):
        prefix = settings.FORCE_SCRIPT_NAME or ""
        self.assertEqual(reverse("admin:users_student_data_sync"), prefix + "/admin/users/student/data-sync/")
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Student.objects.count(), 0)
        self.assertContains(response, f'action="{prefix}/admin/users/student/data-sync/"')
        token = response.context["payload"]
        response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertRedirects(response, prefix + "/admin/users/student/data-sync/", fetch_redirect_response=False)
        self.assertEqual(Student.objects.count(), 1)

    def test_identity_dashboard_preview_and_apply(self):
        user = CustomUser.objects.create_user("old@example.com", None)
        Student.objects.create(user=user, registration_number="OLD", name="Original", email=user.email)
        response = self.upload("identity", f"user,registration_number,email\n{user.pk},NEW,new@example.com\n".encode())
        self.assertContains(response, "IDENTITY UPDATE")
        self.assertContains(response, "Current registration")
        self.assertContains(response, "Identity rows to update")
        self.assertContains(response, "old@example.com")
        self.assertContains(response, "new@example.com")
        self.assertTrue(Student.objects.filter(pk="OLD").exists())
        response = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Student.objects.get(user=user).pk, "NEW")
        user.refresh_from_db()
        self.assertEqual(user.email, "new@example.com")

    def test_identity_requires_superuser_even_with_sync_permissions(self):
        staff = CustomUser.objects.create_user("staff@example.com", None, is_staff=True)
        staff.user_permissions.add(*Permission.objects.filter(codename__in=("change_student", "add_student", "add_customuser")))
        self.client.force_login(staff)
        response = self.upload("identity", b"user,registration_number,email\n1,NEW,new@example.com\n")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.upload().status_code, 200)

    def test_changelist_preserves_import_export_and_sync_link(self):
        response = self.client.get(reverse("admin:users_student_changelist").removeprefix(settings.FORCE_SCRIPT_NAME or ""))
        self.assertContains(response, "Student Data Sync")
        self.assertContains(response, "Import")
        self.assertContains(response, "Export")

    def test_tampered_or_missing_token_cannot_apply(self):
        response = self.client.post(self.url, {"payload": "tampered", "action": "apply"})
        self.assertContains(response, "Preview expired or invalid")
        response = self.client.post(self.url, {"action": "apply"})
        self.assertContains(response, "Validate and preview")
        self.assertEqual(Student.objects.count(), 0)

    def test_apply_revalidates_database(self):
        response = self.upload()
        CustomUser.objects.create_user("NEW@example.com", None, user_type="security")
        response = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply"})
        self.assertContains(response, "protected admin/security")
        self.assertEqual(Student.objects.count(), 0)

    def test_options_change_requires_second_confirmation(self):
        user = CustomUser.objects.create_user("student@example.com", None)
        Student.objects.create(user=user, registration_number="001", name="Name", picture="https://example.com/a.jpg")
        response = self.upload("picture", b"email,picture\nstudent@example.com,\n")
        response = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply", "allow_blank_picture": "on"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(Student.objects.get(pk="001").picture)
        self.client.post(self.url, {"payload": response.context["payload"], "action": "apply", "allow_blank_picture": "on"})
        self.assertIsNone(Student.objects.get(pk="001").picture)

    def test_permission_required(self):
        staff = CustomUser.objects.create_user("staff@example.com", None, is_staff=True)
        staff.user_permissions.add(Permission.objects.get(codename="change_student"))
        self.client.force_login(staff)
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_csrf_required(self):
        from django.test import Client
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        self.assertEqual(client.post(self.url, {"action": "apply"}).status_code, 403)

    def test_preview_has_no_database_writes(self):
        with CaptureQueriesContext(connection) as queries:
            response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(all(q["sql"].startswith("SELECT") for q in queries))

    def test_confirmation_bound_to_admin_and_expiry(self):
        token = self.upload().context["payload"]
        other = CustomUser.objects.create_superuser("otheradmin@example.com", "password")
        self.client.force_login(other)
        self.assertEqual(self.client.post(self.url, {"payload": token, "action": "apply"}).status_code, 403)
        self.client.force_login(self.admin)
        from django.core import signing
        original_loads = signing.loads

        def expire_preview(*args, **kwargs):
            if kwargs.get("salt") == "student-data-sync-v1":
                raise signing.SignatureExpired
            return original_loads(*args, **kwargs)

        with patch("apps.users.student_sync_admin.signing.loads", side_effect=expire_preview):
            response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertContains(response, "Preview expired or invalid")
        self.assertEqual(Student.objects.count(), 0)

    def test_extra_csv_and_xlsx_columns_are_ignored_and_reported(self):
        from openpyxl import Workbook
        from apps.users.services.student_sync_storage import decode_token, read_preview
        hostel = Hostel.objects.create(name="Annual", email="hostel@example.com", contact_number="123")
        headers = ["email", "registration_number", "name", "hostel", "Caretaker Name", "user", "picture"]
        for i, extension in enumerate(("csv", "xlsx")):
            values = [f"annual{i}@example.com", f"A{i}", "Annual Student", "Annual", "Do Not Store", str(self.admin.pk), "https://example.com/photo.jpg"]
            if extension == "csv":
                data = (",".join(headers) + "\n" + ",".join(values) + "\n").encode()
            else:
                workbook = Workbook()
                workbook.active.append(headers)
                workbook.active.append(values)
                output = io.BytesIO()
                workbook.save(output)
                data = output.getvalue()
            with self.subTest(extension=extension):
                response = self.client.post(self.url, {"mode": "full", "file": SimpleUploadedFile(f"annual.{extension}", data)})
                self.assertContains(response, "Ignored columns: Caretaker Name, user")
                self.assertFalse(response.context["plan"].errors)
                token = response.context["payload"]
                stored = read_preview(decode_token(token, self.admin.pk, self.client.session.session_key))
                self.assertNotIn("user", stored["rows"][0])
                self.assertNotIn("caretaker name", stored["rows"][0])
                response = self.client.post(self.url, {"payload": token, "action": "apply"})
                self.assertEqual(response.status_code, 302)
                student = Student.objects.get(pk=f"A{i}")
                self.assertEqual(student.user.email, values[0])
                self.assertNotEqual(student.user_id, self.admin.pk)
                self.assertEqual(student.hostel_id, hostel.pk)
                self.assertEqual(student.picture, values[-1])
                self.assertEqual(student.name, "Annual Student")

    def test_mode_ignored_columns_and_header_errors(self):
        user = CustomUser.objects.create_user("student@example.com", None)
        Student.objects.create(user=user, registration_number="001", name="Preserve")
        response = self.upload("picture", b"email,picture,name,Caretaker Name,user\nstudent@example.com,https://example.com/p.jpg,Ignore,Ignore,1\n")
        self.assertContains(response, "Ignored columns: name, Caretaker Name, user")
        self.client.post(self.url, {"payload": response.context["payload"], "action": "apply"})
        self.assertEqual(Student.objects.get(pk="001").name, "Preserve")
        for data in (
            b"email,name,registration_number,user,USER\na@example.com,A,001,1,2\n",
            b"email,name,registration_number, \na@example.com,A,001,1\n",
        ):
            with self.subTest(data=data):
                response = self.upload(data=data)
                self.assertContains(response, "Headers must be nonblank and unique")

    def test_large_preview_uses_small_token_and_server_file(self):
        from django.core import signing
        from apps.users.services.student_sync_storage import decode_token, read_preview
        generator = random.Random(2026)
        rows = [{"email": f"large{i}@example.com", "registration_number": f"L{i}",
                 "name": "".join(generator.choices(string.ascii_letters, k=90)),
                 "picture": "https://example.com/" + "".join(generator.choices(string.ascii_letters + string.digits, k=165))}
                for i in range(15000)]
        self.assertGreater(len(signing.dumps({"rows": rows}, compress=True)), 1800000)
        data = "email,registration_number,name,picture\n" + "\n".join(",".join(row.values()) for row in rows)
        self.assertLess(len(data.encode()), 10 * 1024 * 1024)
        response = self.upload(data=data.encode())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["plan"].counts["valid_rows"], 15000)
        token = response.context["payload"]
        self.assertLess(len(token), 1000)
        self.assertLess(len(response.content), 30000)
        metadata = decode_token(token, self.admin.pk, self.client.session.session_key)
        self.assertNotIn("rows", metadata)
        self.assertEqual(read_preview(metadata)["rows"], rows)
        self.assertEqual(Student.objects.count(), 0)
        self.assertNotIn("rows", self.client.session)

    def test_success_cleans_temporary_data_and_prevents_replay(self):
        token = self.upload().context["payload"]
        self.assertEqual(len(list(self.storage.glob("*.json"))), 1)
        response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(list(self.storage.iterdir()), [])
        response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertContains(response, "Preview expired or invalid")
        self.assertEqual(Student.objects.count(), 1)

    def test_real_expiry_and_tampering_leave_database_unchanged(self):
        import time
        token = self.upload().context["payload"]
        response = self.client.post(self.url, {"payload": token + "x", "action": "apply"})
        self.assertContains(response, "Preview expired or invalid")
        with patch("apps.users.services.student_sync_storage.time.time", return_value=time.time() + 1801):
            response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertContains(response, "Preview expired or invalid")
        self.assertEqual(Student.objects.count(), 0)

    def test_failed_revalidation_restores_preview_for_retry(self):
        token = self.upload().context["payload"]
        conflict = CustomUser.objects.create_user("NEW@example.com", None, user_type="security")
        response = self.client.post(self.url, {"payload": token, "action": "apply"})
        self.assertContains(response, "protected admin/security")
        self.assertEqual(len(list(self.storage.glob("*.json"))), 1)
        self.assertEqual(list(self.storage.glob("*.applying")), [])
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(CustomUser.objects.get(pk=conflict.pk).user_type, "security")

    def test_token_bound_to_login_session(self):
        token = self.upload().context["payload"]
        self.client.logout()
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(self.url, {"payload": token, "action": "apply"}).status_code, 403)

    def real_picture_file(self, picture):
        return (self.REAL_HEADERS + "\n001,Updated,Annual,Caretaker,F,202,123,real@example.com,456,2,"
                + str(self.admin.pk) + "," + picture + "\n").encode()

    def create_real_picture_student(self):
        hostel = Hostel.objects.create(name="Annual", email="hostel@example.com", contact_number="123")
        user = CustomUser.objects.create_user("real@example.com", None)
        return Student.objects.create(user=user, registration_number="001", name="Original", hostel=hostel,
                                      picture="https://example.com/original.jpg")

    def test_real_url_header_updates_picture_in_both_modes_and_formats(self):
        from openpyxl import Workbook
        from apps.users.services.student_sync_storage import decode_token, read_preview
        student = self.create_real_picture_student()
        for mode in ("picture", "full"):
            for extension in ("csv", "xlsx"):
                with self.subTest(mode=mode, extension=extension):
                    before = Student.objects.values().get(pk=student.pk)
                    picture = f"https://example.com/{mode}-{extension}.jpg"
                    data = self.real_picture_file(picture)
                    if extension == "xlsx":
                        book = Workbook()
                        for line in data.decode().splitlines():
                            book.active.append(line.split(","))
                        output = io.BytesIO()
                        book.save(output)
                        data = output.getvalue()
                    response = self.client.post(self.url, {"mode": mode, "file": SimpleUploadedFile(f"real.{extension}", data)})
                    self.assertContains(response, "Mapped picture column: URL → picture")
                    plan = response.context["plan"]
                    self.assertFalse(plan.errors)
                    self.assertNotIn("URL", plan.ignored_columns)
                    self.assertIn("Caretaker Name", plan.ignored_columns)
                    self.assertIn("user", plan.ignored_columns)
                    token = response.context["payload"]
                    stored = read_preview(decode_token(token, self.admin.pk, self.client.session.session_key))
                    self.assertEqual(stored["rows"][0]["picture"], picture)
                    self.assertNotIn("url", stored["rows"][0])
                    self.assertNotIn("user", stored["rows"][0])
                    self.assertNotIn("caretaker name", stored["rows"][0])
                    result = self.client.post(self.url, {"payload": token, "action": "apply"})
                    self.assertEqual(result.status_code, 302)
                    after = Student.objects.values().get(pk=student.pk)
                    self.assertEqual(after["picture"], picture)
                    self.assertEqual(after["user_id"], before["user_id"])
                    if mode == "picture":
                        before["picture"] = picture
                        self.assertEqual(after, before)
                    else:
                        self.assertEqual(after["name"], "Updated")

    def test_real_blank_url_preserves_until_explicit_picture_clear(self):
        student = self.create_real_picture_student()
        for mode in ("full", "picture"):
            response = self.upload(mode, self.real_picture_file(""))
            self.assertFalse(response.context["plan"].errors)
            token = response.context["payload"]
            self.assertEqual(self.client.post(self.url, {"payload": token, "action": "apply"}).status_code, 302)
            student.refresh_from_db()
            self.assertEqual(student.picture, "https://example.com/original.jpg")
        response = self.upload("picture", self.real_picture_file(""))
        response = self.client.post(self.url, {"payload": response.context["payload"], "action": "preview", "allow_blank_picture": "on"})
        self.assertEqual(response.context["plan"].counts["pictures_updated"], 1)
        result = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply", "allow_blank_picture": "on"})
        self.assertEqual(result.status_code, 302)
        student.refresh_from_db()
        self.assertIsNone(student.picture)

    def test_real_headers_with_picture_and_url_are_rejected(self):
        student = self.create_real_picture_student()
        lines = self.real_picture_file("https://example.com/url.jpg").decode().splitlines()
        data = (lines[0] + ",picture\n" + lines[1] + ",https://example.com/picture.jpg\n").encode()
        for mode in ("full", "picture"):
            response = self.upload(mode, data)
            self.assertContains(response, "Ambiguous picture columns: URL, picture")
            self.assertNotIn("payload", response.context)
        student.refresh_from_db()
        self.assertEqual(student.picture, "https://example.com/original.jpg")
        self.assertEqual(list(self.storage.iterdir()), [])


class StudentSyncPictureHeaderTests(SimpleTestCase):
    def test_all_aliases_are_case_insensitive_and_normalized(self):
        from openpyxl import Workbook
        for alias in ("picture", "URL", "PiCtUrE_uRl", "ImAgE_uRl"):
            for extension in ("csv", "xlsx"):
                with self.subTest(alias=alias, extension=extension):
                    if extension == "csv":
                        data = f"email,{alias}\nstudent@example.com,https://example.com/p.jpg\n".encode()
                    else:
                        book = Workbook()
                        book.active.append(["email", alias])
                        book.active.append(["student@example.com", "https://example.com/p.jpg"])
                        output = io.BytesIO()
                        book.save(output)
                        data = output.getvalue()
                    headers, rows = read_upload(SimpleUploadedFile(f"pictures.{extension}", data))
                    self.assertEqual(headers, ["email", alias])
                    self.assertEqual(rows, [{"email": "student@example.com", "picture": "https://example.com/p.jpg"}])

    def test_every_duplicate_picture_alias_combination_is_ambiguous(self):
        from itertools import combinations_with_replacement
        from apps.users.services.student_sync import normalize_headers
        for first, second in combinations_with_replacement(("picture", "URL", "picture_url", "image_url"), 2):
            with self.subTest(first=first, second=second), self.assertRaisesMessage(ValidationError, "Ambiguous picture columns"):
                normalize_headers(["email", first, second.upper()])

class StudentSyncRowHandlingTests(TestCase):
    """Exercise raw uploads, stored previews and apply in both formats/modes."""

    setUp = StudentSyncAdminTests.setUp
    create_real_picture_student = StudentSyncAdminTests.create_real_picture_student

    def row_upload(self, mode, extension, rows):
        headers = ["email", "registration_number", "name", "URL", "gender"]
        output = io.BytesIO()
        if extension == "xlsx":
            from openpyxl import Workbook
            book = Workbook()
            for row in [headers] + rows:
                book.active.append(row)
            book.save(output)
            data = output.getvalue()
        else:
            import csv
            stream = io.StringIO()
            csv.writer(stream).writerows([headers] + rows)
            data = stream.getvalue().encode()
        return self.client.post(self.url, {"mode": mode, "file": SimpleUploadedFile(f"annual.{extension}", data)})

    def test_raw_identity_conflicts_and_first_physical_row(self):
        student = self.create_real_picture_student()
        for mode in ("picture", "full"):
            for extension in ("csv", "xlsx"):
                for identity in (("001", ""), ("", "Student")):
                    response = self.row_upload(mode, extension, [["", *identity, "2", ""]])
                    self.assertEqual(response.context["plan"].counts["error_rows"], 1)
                    self.assertContains(response, "email is required")
                first = ["real@example.com", "001", "Original", "https://example.com/a.jpg", "M"]
                second = [" REAL@EXAMPLE.COM ", "001", "Original", "https://example.com/b.jpg", "male"]
                response = self.row_upload(mode, extension, [["", "", "", "2", ""], first, second])
                plan = response.context["plan"]
                self.assertEqual(plan.counts["duplicate_emails"], 1)
                self.assertEqual(plan.counts["error_rows"], 2)
                self.assertContains(response, "rows 3, 4")
                result = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply"})
                self.assertContains(result, "conflicting duplicate email")
                student.refresh_from_db()
                self.assertEqual(student.picture, "https://example.com/original.jpg")
                # Invalid identical rows still validate the first physical occurrence.
                first[3] = second[3] = "invalid"
                response = self.row_upload(mode, extension, [["", "", "", "2", ""], first, second])
                self.assertEqual(response.context["plan"].counts["identical_duplicate_rows_skipped"], 1)
                self.assertEqual(response.context["plan"].counts["error_rows"], 1)
                self.assertTrue(all(error.startswith("Row 3:") for error in response.context["plan"].errors))

    def test_13000_annual_rows_with_duplicate_and_footer(self):
        for extension in ("csv", "xlsx"):
            for mode in ("full", "picture"):
                with self.subTest(mode=mode, extension=extension):
                    picture = f"https://example.com/{extension}-{mode}.jpg"
                    rows = [[f"annual{i}@example.com", f"R{i:05}", f"Student {i}", picture, "M"] for i in range(13000)]
                    rows.append([" ANNUAL0@EXAMPLE.COM ", "R00000", "Student 0", picture, "male"])
                    if mode == "picture":
                        rows[-1][1:3] = ["Ignored registration", "Ignored name"]
                    rows.append(["", "", "", "2", ""])
                    response = self.row_upload(mode, extension, rows)
                    plan = response.context["plan"]
                    self.assertFalse(plan.errors)
                    for key, value in {"total_rows": 13002, "rows_considered": 13000, "valid_rows": 13000,
                                       "identical_duplicate_rows_skipped": 1, "footer_rows_skipped": 1,
                                       "duplicate_emails": 0, "error_rows": 0, "unchanged_rows": 0}.items():
                        self.assertEqual(plan.counts[key], value, key)
                    for label in ("Total physical data rows", "Rows considered for sync", "Unchanged rows",
                                  "Identical duplicate rows skipped", "Non-student/footer rows skipped", "Conflicting duplicate emails"):
                        self.assertContains(response, label)
                    result = self.client.post(self.url, {"payload": response.context["payload"], "action": "apply"})
                    self.assertEqual(result.status_code, 302)
                    self.assertEqual(Student.objects.count(), 13000)
                    self.assertEqual(CustomUser.objects.filter(user_type="student").count(), 13000)
                    self.assertEqual(Student.objects.filter(picture=picture, gender="male").count(), 13000)
                    repeat = self.row_upload(mode, extension, rows).context["plan"]
                    self.assertFalse(repeat.errors)
                    self.assertEqual(repeat.counts["unchanged_rows"], 13000)
                    self.assertEqual(repeat.counts["students_to_update"], 0)

    def test_identical_duplicates_do_not_bypass_protection_or_unknown_hostels(self):
        for row in (
            {"email": self.admin.email, "registration_number": "P", "name": "Protected"},
            {"email": "new@example.com", "registration_number": "N", "name": "New", "hostel": "Unknown"},
        ):
            plan = preview_sync(list(row), [row, row.copy()], "full")
            self.assertEqual(plan.counts["identical_duplicate_rows_skipped"], 1)
            self.assertEqual(plan.counts["error_rows"], 1)
            with self.assertRaises(ValidationError):
                apply_sync(list(row), [row, row.copy()], "full", actor=self.admin.pk, filename="test.csv")
        self.assertEqual(Student.objects.count(), 0)
